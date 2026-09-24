"""Editable four-panel method slide. Output: figures/fig_arch.pptx."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent / "figures" / "fig_arch.pptx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x3D, 0x3D, 0x3D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE_BG = RGBColor(0xE7, 0xF0, 0xFB)
BLUE_LN = RGBColor(0x3D, 0x6E, 0xA8)
YEL_BG = RGBColor(0xFF, 0xF6, 0xDE)
YEL_LN = RGBColor(0xC4, 0x96, 0x3A)
GRN_BG = RGBColor(0xE5, 0xF6, 0xEC)
GRN_LN = RGBColor(0x3E, 0x8F, 0x62)
PNK_BG = RGBColor(0xFD, 0xEC, 0xEC)
PNK_LN = RGBColor(0xC4, 0x6A, 0x6A)
POS = RGBColor(0xE3, 0xF4, 0xDC)
POS_LN = RGBColor(0x6A, 0xA8, 0x4E)
NEG = RGBColor(0xFA, 0xE2, 0xE2)
NEG_LN = RGBColor(0xC4, 0x78, 0x78)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE = RGBColor(0xEE, 0xE4, 0xF8)
PURPLE_LN = RGBColor(0x6E, 0x56, 0xA0)
GOLD = RGBColor(0xE8, 0xF7, 0xEF)
GOLD_LN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xF8, 0xDC, 0xDC)
RED_LN = RGBColor(0xB0, 0x40, 0x40)


def _fill_line(shape, fill, line, lw=1.0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(lw)


def _no_margin(tf):
    tf.word_wrap = True
    tf.auto_size = None
    body = tf._txBody
    bodyPr = body.find(qn("a:bodyPr"))
    if bodyPr is not None:
        for attr in ("lIns", "tIns", "rIns", "bIns"):
            bodyPr.set(attr, str(Emu(Inches(0.06))))


def card(slide, l, t, w, h, fill, line, lines, size=11, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, lw=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.adjustments[0] = 0.08
    _fill_line(shape, fill, line, lw)
    tf = shape.text_frame
    tf.clear()
    _no_margin(tf)
    tf.word_wrap = True
    tf.auto_size = None
    shape.text_frame.paragraphs[0].alignment = align
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        if isinstance(line, tuple):
            text, sz, is_bold, color = line
        else:
            text, sz, is_bold, color = line, size, bold, INK
        run = p.add_run()
        run.text = text
        run.font.size = Pt(sz)
        run.font.bold = is_bold
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return shape


def arrow(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def down_arrow(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()

    # panel shells
    card(slide, 0.08, 0.08, 6.52, 3.62, BLUE_BG, BLUE_LN, [], lw=1.25)
    card(slide, 6.74, 0.08, 6.52, 3.62, YEL_BG, YEL_LN, [], lw=1.25)
    card(slide, 0.08, 3.82, 6.52, 3.60, GRN_BG, GRN_LN, [], lw=1.25)
    card(slide, 6.74, 3.82, 6.52, 3.60, PNK_BG, PNK_LN, [], lw=1.25)

    # ----- 1 -----
    card(slide, 0.18, 0.14, 6.32, 0.32, BLUE_BG, BLUE_BG, [
        ("1.  Build a domain steering vector (CAA)", 14, True, INK),
    ], lw=0)
    card(slide, 0.18, 0.48, 2.72, 1.72, POS, POS_LN, [
        ("Expert  (+)", 12, True, RGBColor(0x1E, 0x6B, 0x32)),
        ("You are an Electrical engineering practitioner.", 10, False, MUTED),
        ("How does feedback work?", 11, False, INK),
        ("If the feedback is negative, does it", 11, False, INK),
        ("make things worse?", 11, False, INK),
        ("… (28 more)     30 pairs", 11, True, INK),
    ])
    card(slide, 0.18, 2.28, 2.72, 1.28, NEG, NEG_LN, [
        ("Default  (−)", 12, True, RGBColor(0xA0, 0x30, 0x30)),
        ("You are a helpful assistant.", 10, False, MUTED),
        ("How does feedback work?", 11, False, INK),
        ("If the feedback is negative, …", 11, False, INK),
        ("Same stems.  … (28 more)", 11, True, INK),
    ])
    arrow(slide, 2.96, 1.15, 0.28, 0.18, BLUE_LN)
    card(slide, 3.28, 0.48, 1.48, 1.72, CARD, BLUE_LN, [
        ("Mean-pool", 12, True, INK),
        ("assistant tokens", 11, False, INK),
        ("layers 10–20", 12, True, INK),
        ("Replayed under", 10, False, MUTED),
        ("the helper", 10, False, MUTED),
        ("template.", 10, False, MUTED),
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    arrow(slide, 4.82, 1.15, 0.26, 0.18, BLUE_LN)
    card(slide, 5.12, 0.48, 1.36, 1.72, CARD, BLUE_LN, [
        ("v = unit(μ+ − μ−)", 11, True, INK),
        ("one unit vector", 10, False, MUTED),
        ("at each of", 10, False, MUTED),
        ("layers 10–20", 12, True, INK),
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    card(slide, 3.28, 2.32, 3.20, 1.22, CARD, BLUE_LN, [
        ("Same unnamed stem on both sides.", 12, True, INK),
        ("The domain prompt only writes the expert text. It is not in the states that are averaged.", 11, False, INK),
    ], anchor=MSO_ANCHOR.MIDDLE)

    # ----- 2 -----
    card(slide, 6.84, 0.14, 6.32, 0.32, YEL_BG, YEL_BG, [
        ("2.  Add the vector at generation time", 14, True, INK),
    ], lw=0)
    card(slide, 6.90, 0.52, 3.55, 1.15, CARD, YEL_LN, [
        ("Input", 12, True, INK),
        ("System prompt: helper, or", 11, False, INK),
        ("“You are an Electrical engineering", 11, False, INK),
        ("practitioner.”", 11, False, INK),
        ("Question does not name the domain.", 11, True, INK),
    ])
    card(slide, 6.90, 1.78, 3.55, 1.72, CARD, YEL_LN, [
        ("Layer 1", 12, False, MUTED),
        ("⋮", 12, False, MUTED),
        ("one layer in 10–20   (add v)", 13, True, INK),
        ("every token, including system", 11, False, INK),
        ("and user turns", 11, False, INK),
        ("⋮", 12, False, MUTED),
        ("Generate answer", 12, True, INK),
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    card(slide, 10.55, 0.70, 2.52, 2.55, CARD, YEL_LN, [
        ("h  ←  h + α‖h‖v", 14, True, INK),
        ("h   residual stream", 11, False, MUTED),
        ("v   domain steering vector", 11, False, MUTED),
        ("α   strength, a fraction", 11, False, MUTED),
        ("of ‖h‖ at that token", 11, False, MUTED),
    ], anchor=MSO_ANCHOR.MIDDLE)

    # ----- 3 -----
    card(slide, 0.18, 3.88, 6.32, 0.30, GRN_BG, GRN_BG, [
        ("3.  Evaluate on polysemy traps", 14, True, INK),
    ], lw=0)
    card(slide, 0.18, 4.22, 2.15, 2.87, CARD, GRN_LN, [
        ("Question", 12, True, INK),
        ("no domain name", 10, False, MUTED),
        ("How much can a cell hold?", 13, True, INK),
        ("Electrical engineering", 11, False, MUTED),
        ("15 items × 144 domains", 10, False, MUTED),
    ], anchor=MSO_ANCHOR.MIDDLE)
    arms = [
        (4.22, "Unsteered", "helper, vector off", "A typical cell can hold a maximum of about 1.5 liters of blood…"),
        (4.96, "Prompt", "practitioner prompt, vector off", "A typical lithium-ion cell can hold around 3.7 volts and 2–3 ampere-hours (Ah) of charge."),
        (5.70, "CAA", "helper, vector on", "A cell can hold 1 coulomb of charge."),
        (6.44, "prompt+CAA", "practitioner prompt, vector on", "The amount of charge a cell can hold is measured in ampere-hours (Ah) or milliampere-hours (mAh)."),
    ]
    for top, name, spec, ans in arms:
        fill, line = (PURPLE, PURPLE_LN) if "CAA" in name else (CARD, GRN_LN)
        card(slide, 2.42, top, 4.02, 0.70, fill, line, [
            (f"{name}   {spec}", 10, True, INK),
            (ans, 10, False, MUTED),
        ], anchor=MSO_ANCHOR.MIDDLE)

    # ----- 4 -----
    card(slide, 6.84, 3.88, 6.32, 0.30, PNK_BG, PNK_BG, [
        ("4.  Score and decide", 14, True, INK),
    ], lw=0)
    card(slide, 6.90, 4.26, 2.55, 2.95, CARD, PNK_LN, [
        ("MPNet cosine", 13, True, INK),
        ("all-mpnet-base-v2", 11, False, MUTED),
        ("to one frozen reference", 12, True, INK),
        ("", 6, False, INK),
        ("Reference (domain sense)", 11, True, RGBColor(0x1E, 0x6B, 0x32)),
        ("Capacity specifies how much", 12, False, INK),
        ("charge the cell can supply.", 12, False, INK),
        ("", 6, False, INK),
        ("Every arm is scored against", 11, False, MUTED),
        ("this same sentence.", 11, False, MUTED),
    ])
    card(slide, 9.55, 4.26, 3.52, 0.42, CARD, PNK_LN, [
        ("Pick the closer answer", 13, True, INK),
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    card(slide, 9.55, 4.78, 3.52, 1.05, GOLD, GOLD_LN, [
        ("Item win", 12, True, RGBColor(0x1E, 0x6B, 0x32)),
        ("Closer to the reference than", 12, False, INK),
        ("the other arm. Ties lose.", 12, True, INK),
    ], anchor=MSO_ANCHOR.MIDDLE)
    card(slide, 9.55, 5.95, 3.52, 1.22, RED, RED_LN, [
        ("Domain majority", 12, True, RGBColor(0x8C, 0x28, 0x28)),
        ("At least 8 of 15 answers", 13, True, INK),
        ("are closer.", 13, True, INK),
    ], anchor=MSO_ANCHOR.MIDDLE)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
